import streamlit as st
from docx import Document
from docx.shared import Pt
from io import BytesIO
import os
import pandas as pd
from datetime import date

st.set_page_config(page_title="Protokol Generatoru", layout="wide")
st.title("📋 Qiymət Razılaşma Protokolu")

# --- ŞABLONUN YERİ ---
current_dir = os.path.dirname(os.path.abspath(__file__))
sablon_yolu = os.path.join(current_dir, "sablon_protokol.docx")

if not os.path.exists(sablon_yolu):
    st.error("❌ 'sablon_protokol.docx' faylı tapılmadı!")
    st.stop()

# --- FORM HİSSƏSİ ---
col1, col2 = st.columns([1, 2])

with col1:
    st.subheader("1. Əsas Məlumatlar")
    nomre = st.text_input("Protokol №", value="1")
    tarix_val = st.date_input("Tarix", value=date.today())
    tarix = tarix_val.strftime("%d.%m.%Y")
    
    # --- DƏYİŞİKLİK BURADADIR (BOŞ XANALAR) ---
    sirket = st.text_input("Müştəri Şirkət", value="", placeholder="Şirkətin adı...")
    musteri = st.text_input("Direktor", value="", placeholder="Ad Soyad")
    
    st.subheader("2. Bank Rekvizitləri")
    voen = st.text_input("VÖEN")
    hesab = st.text_input("Hesab (H/h)")
    bank = st.text_input("Bank Adı")
    bank_voen = st.text_input("Bank VÖEN")
    mh = st.text_input("M/h")
    swift = st.text_input("SWIFT")
    kod = st.text_input("Kod")

with col2:
    st.subheader("3. Məhsul Siyahısı")
    
    # Boş cədvəl strukturu
    df = pd.DataFrame(columns=["Malın/Xidmətin Adı", "Ölçü Vahidi", "Miqdar", "Qiymət (AZN)"])

    edited_df = st.data_editor(
        df,
        num_rows="dynamic",
        use_container_width=True,
        column_config={
            "Malın/Xidmətin Adı": st.column_config.TextColumn(
                "Malın/Xidmətin Adı",
                width="large",
                required=True,
                help="Məhsulun adını bura yazın"
            ),
            "Ölçü Vahidi": st.column_config.SelectboxColumn(
                "Ölçü Vahidi",
                width="medium",
                options=[
                    "ədəd", 
                    "metr", 
                    "kq", 
                    "litr", 
                    "saat", 
                    "gün", 
                    "ay", 
                    "dəst", 
                    "qutu", 
                    "komplekt",
                    "xidmət"
                ],
                required=True
            ),
            "Miqdar": st.column_config.NumberColumn(
                "Miqdar",
                min_value=0,
                step=1,
                format="%d"
            ),
            "Qiymət (AZN)": st.column_config.NumberColumn(
                "Qiymət (AZN)",
                min_value=0,
                format="%.2f ₼"
            )
        },
        key="editor"
    )

    # --- HESABLAMA ---
    if not edited_df.empty:
        # None dəyərləri 0-a çeviririk
        edited_df["Miqdar"] = pd.to_numeric(edited_df["Miqdar"], errors='coerce').fillna(0)
        edited_df["Qiymət (AZN)"] = pd.to_numeric(edited_df["Qiymət (AZN)"], errors='coerce').fillna(0)
        
        edited_df["Məbləğ"] = edited_df["Miqdar"] * edited_df["Qiymət (AZN)"]
        cem = edited_df["Məbləğ"].sum()
    else:
        cem = 0.0

    edv = cem * 0.18
    yekun = cem + edv

    st.divider()
    c1, c2, c3 = st.columns(3)
    c1.metric("Cəmi", f"{cem:.2f} ₼")
    c2.metric("ƏDV (18%)", f"{edv:.2f} ₼")
    c3.metric("YEKUN", f"{yekun:.2f} ₼")

# --- HAZIRLAMAQ ---
if st.button("Protokolu Hazırla", type="primary"):
    doc = Document(sablon_yolu)

    # Şrift tənzimləməsi (Calibri)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    deyisenler = {
        "{{NOMRE}}": nomre,
        "{{TARIX}}": tarix,
        "{{SIRKET}}": sirket,
        "{{MUSTERI}}": musteri,
        "{{VOEN}}": voen,
        "{{HESAB}}": hesab,
        "{{BANK}}": bank,
        "{{BANK_VOEN}}": bank_voen,
        "{{MH}}": mh,
        "{{SWIFT}}": swift,
        "{{KOD}}": kod,
    }

    # Mətn Dəyişdirmə Funksiyası
    def replace_text_bold(paragraph):
        for key, value in deyisenler.items():
            if key in paragraph.text:
                if value:
                    if key in paragraph.text: 
                        parts = paragraph.text.split(key)
                        paragraph.clear()
                        for i in range(len(parts) - 1):
                            run = paragraph.add_run(parts[i])
                            run.font.name = 'Calibri'
                            
                            run = paragraph.add_run(str(value))
                            run.bold = True
                            run.font.name = 'Calibri'
                        
                        run = paragraph.add_run(parts[-1])
                        run.font.name = 'Calibri'
                else:
                    paragraph.text = paragraph.text.replace(key, "")

    # Mətni yenilə
    for p in doc.paragraphs:
        replace_text_bold(p)

    # Cədvəlləri yenilə
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_text_bold(p)

    # 3. Məhsul Cədvəlini Word-ə Yazmaq
    try:
        table = doc.tables[0] 
        
        # Cədvəli təmizlə
        for i in range(len(table.rows) - 1, 0, -1):
            row = table.rows[i]
            tr = row._element
            tr.getparent().remove(tr)

        # Məhsulları əlavə et
        for index, row in edited_df.iterrows():
            new_row = table.add_row()
            
            mal_adi = str(row["Malın/Xidmətin Adı"]) if row["Malın/Xidmətin Adı"] else ""
            olcu = str(row["Ölçü Vahidi"]) if row["Ölçü Vahidi"] else ""
            
            new_row.cells[0].text = str(index + 1)
            new_row.cells[1].text = mal_adi
            new_row.cells[2].text = olcu
            new_row.cells[3].text = str(int(row["Miqdar"])) # Tam ədəd kimi göstər
            new_row.cells[4].text = f"{row['Qiymət (AZN)']:.2f}"
            new_row.cells[5].text = f"{row['Məbləğ']:.2f}"
            
            for cell in new_row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Calibri'

        # Yekunları əlavə et
        def add_summary_row(label, value):
            row = table.add_row()
            row.cells[4].text = label
            row.cells[5].text = value
            for cell in [row.cells[4], row.cells[5]]:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = 'Calibri'
                        r.bold = True

        add_summary_row("Cəmi", f"{cem:.2f}")
        add_summary_row("ƏDV 18%", f"{edv:.2f}")
        add_summary_row("Yekun", f"{yekun:.2f}")

    except Exception as e:
        st.error(f"Cədvəl xətası: {e}")

    bio = BytesIO()
    doc.save(bio)
    
    st.success("✅ Protokol hazırdır!")
    st.download_button(
        label="📥 Protokolu Yüklə",
        data=bio.getvalue(),
        file_name=f"Protokol_{sirket}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
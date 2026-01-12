import streamlit as st
from docx import Document
from io import BytesIO
import os
from datetime import date, timedelta

# Səhifənin parametrləri
st.set_page_config(page_title="Müqavilə Generatoru", layout="centered")
st.title("📄 Avtomatik Müqavilə Hazırlayan")

# --- ŞABLON FAYLINI TAPMAQ ---
current_dir = os.path.dirname(os.path.abspath(__file__))
sablon_yolu = os.path.join(current_dir, "sablon.docx")

if not os.path.exists(sablon_yolu):
    st.error(f"❌ XƏTA: 'sablon.docx' faylı tapılmadı!")
    st.warning(f"Zəhmət olmasa 'sablon.docx' faylını bu qovluğa qoyun: {current_dir}")
    st.stop()

# --- FORM HİSSƏSİ ---
with st.form("muqavile_formu"):
    st.subheader("1. Əsas Məlumatlar")
    col1, col2 = st.columns(2)
    
    with col1:
        nomre = st.text_input("Müqavilə Nömrəsi", placeholder="məs: 055")
        
        # --- TARİX SEÇİMİ (TƏQVİM) ---
        # 1. İmzalanma tarixi (Varsayılan: Bu gün)
        secilen_tarix = st.date_input("İmzalanma Tarixi", value=date.today())
        
        # 2. Bitmə tarixi (Varsayılan: 1 il sonra)
        secilen_bitme = st.date_input("Bitmə Tarixi", value=date.today() + timedelta(days=365))
        
        # Tarixləri "Gün.Ay.İl" formatına çeviririk (String)
        tarix_str = secilen_tarix.strftime("%d.%m.%Y")
        bitme_str = secilen_bitme.strftime("%d.%m.%Y")
    
    with col2:
        sirket = st.text_input("Şirkət Adı", placeholder="məs: 'ABC' MMC")
        musteri = st.text_input("Direktorun Adı", placeholder="məs: Əliyev Vəli")

    st.subheader("2. Bank Rekvizitləri")
    col3, col4 = st.columns(2)
    
    with col3:
        voen = st.text_input("Şirkətin VÖEN-i")
        hesab = st.text_input("Hesab Nömrəsi (H/h)")
        mh = st.text_input("M/h (Müxbir Hesab)")
        
    with col4:
        bank = st.text_input("Bankın Adı")
        bank_voen = st.text_input("Bankın VÖEN-i")
        swift = st.text_input("SWIFT Kodu")
        bank_kodu = st.text_input("Bank Kodu (varsa)")

    submitted = st.form_submit_button("Müqaviləni Hazırla")

# --- KODLAMA HİSSƏSİ ---
if submitted:
    if not sirket or not musteri:
        st.warning("⚠️ Zəhmət olmasa ən azı Şirkət və Direktor adını yazın.")
    else:
        doc = Document(sablon_yolu)
        
        # Dəyişənlər (Formatlanmış tarixləri bura ötürürük)
        deyisenler = {
            "{{NOMRE}}": nomre,
            "{{TARIX}}": tarix_str,        # Avtomatik formatlanmış: 11.01.2026
            "{{BITME_TARIXI}}": bitme_str, # Avtomatik formatlanmış
            "{{SIRKET}}": sirket,
            "{{MUSTERI}}": musteri,
            "{{VOEN}}": voen,
            "{{HESAB}}": hesab,
            "{{BANK}}": bank,
            "{{BANK_VOEN}}": bank_voen,
            "{{MH}}": mh,
            "{{SWIFT}}": swift,
            "{{KOD}}": bank_kodu
        }

        # --- QALIN (BOLD) YAZMA FUNKSİYASI ---
        def replace_text_bold(paragraph):
            for key, value in deyisenler.items():
                if key in paragraph.text:
                    if value: 
                        parts = paragraph.text.split(key)
                        paragraph.clear() 
                        
                        for i in range(len(parts) - 1):
                            paragraph.add_run(parts[i]) 
                            run = paragraph.add_run(value) 
                            run.bold = True # <-- Qalın edir
                        paragraph.add_run(parts[-1]) 
                    else:
                        paragraph.text = paragraph.text.replace(key, " ")

        # 1. Mətni yoxlayırıq
        for p in doc.paragraphs:
            replace_text_bold(p)

        # 2. Cədvəlləri yoxlayırıq
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        replace_text_bold(p)

        # Yaddaşa yazırıq
        bio = BytesIO()
        doc.save(bio)
        
        st.success(f"✅ Müqavilə hazırdır! Tarixlər: {tarix_str} - {bitme_str}")
        
        st.download_button(
            label="📥 Hazır Müqaviləni Yüklə",
            data=bio.getvalue(),
            file_name=f"Muqavile_{sirket}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
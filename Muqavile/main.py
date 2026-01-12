import streamlit as st
from docx import Document
from docx.shared import Pt
from io import BytesIO
import os
import pandas as pd
from datetime import date
import requests
from bs4 import BeautifulSoup
from PIL import Image, ImageDraw, ImageFont
import re

# --- SƏHİFƏ TƏNZİMLƏMƏLƏRİ ---
st.set_page_config(page_title="Sənəd Portalı", layout="wide", page_icon="🗂️")

# --- SESSION STATE (Etiket Siyahısı üçün) ---
if "print_queue" not in st.session_state:
    st.session_state.print_queue = []
if "generated_image" not in st.session_state:
    st.session_state.generated_image = None
if "specs" not in st.session_state:
    st.session_state.specs = {
        "Brand": "", "Model": "", "PartNumber": "", 
        "CPU": "", "RAM": "", "SSD": "", "GPU": "", 
        "Screen": "", "OS": "Windows 11", 
        "PriceOld": "", "PriceNew": "" 
    }

# --- CSS İLƏ PEŞƏKAR FOOTER ---
hide_st_style = """
            <style>
            #MainMenu {visibility: hidden;}
            footer {visibility: hidden;}
            header {visibility: hidden;}
            
            .footer {
                position: fixed;
                left: 0;
                bottom: 0;
                width: 100%;
                background-color: #f8f9fa;
                color: #6c757d;
                text-align: center;
                padding: 10px;
                font-size: 14px;
                border-top: 1px solid #e9ecef;
                z-index: 100;
            }
            </style>
            <div class="footer">
                <p>Bu proqram təminatı <b>Laptop Market</b> şirkətinin mülkiyyətidir. © 2026 Bütün hüquqlar qorunur.</p>
            </div>
            """
st.markdown(hide_st_style, unsafe_allow_html=True)

# ---------------------------------------------------------
# KÖMƏKÇİ FUNKSİYALAR (ETİKET & LOGO)
# ---------------------------------------------------------
def get_absolute_path(relative_path):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_dir, relative_path)

def logo_goster():
    if os.path.exists("logo.png"):
        return st.image("logo.png", width=150)
    return None

def load_font(filename, size):
    # 1. Yerli 'fonts' qovluğu
    local_path = get_absolute_path(os.path.join("fonts", filename))
    if os.path.exists(local_path):
        try: return ImageFont.truetype(local_path, size)
        except: pass
    # 2. Windows Fonts
    windows_path = os.path.join(r"C:\Windows\Fonts", filename)
    if os.path.exists(windows_path):
        try: return ImageFont.truetype(windows_path, size)
        except: pass
    try: return ImageFont.truetype(filename, size)
    except: return ImageFont.load_default()

def make_white(image):
    image = image.convert("RGBA")
    white_img = Image.new("RGBA", image.size, (255, 255, 255, 255))
    final_img = Image.composite(white_img, Image.new("RGBA", image.size, (0,0,0,0)), image)
    return final_img

def clean_price_text(text):
    if not text: return ""
    return re.sub(r'[^\d]', '', text)

def scrape_smart(url):
    headers = {"User-Agent": "Mozilla/5.0"}
    data = {
        "Brand": "", "Model": "", "PartNumber": "", 
        "CPU": "", "RAM": "", "SSD": "", "GPU": "", 
        "Screen": "", "OS": "Windows 11", 
        "PriceOld": "", "PriceNew": ""
    }
    try:
        response = requests.get(url, headers=headers, timeout=15)
        soup = BeautifulSoup(response.text, 'html.parser')
        text = soup.get_text(separator="\n")

        title_tag = soup.find("h1")
        if title_tag:
            title_text = title_tag.get_text(strip=True)
            parts = title_text.split(" ", 2)
            if len(parts) > 0: data["Brand"] = parts[0].strip()
            if len(parts) > 1: data["Model"] = parts[1].strip()
            if len(parts) > 2: data["PartNumber"] = parts[2].strip()

        price_old_tag = soup.find(class_=lambda x: x and ("price-old" in x or "old-price" in x))
        if price_old_tag: data["PriceOld"] = clean_price_text(price_old_tag.get_text())
            
        price_new_tag = soup.find(class_=lambda x: x and ("price-new" in x or "special-price" in x))
        if price_new_tag:
            data["PriceNew"] = clean_price_text(price_new_tag.get_text())
        else:
            price_normal_tag = soup.find(class_="product-price")
            if price_normal_tag: data["PriceNew"] = clean_price_text(price_normal_tag.get_text())

        regexes = {
            "CPU": r'(?:CPU|Prosessor)\s*:\s*([^\n\r]+)',
            "RAM": r'(?:RAM|Operativ)\s*:\s*([^\n\r]+)',
            "SSD": r'(?:SSD|Yaddaş)\s*:\s*([^\n\r]+)',
            "Screen": r'(?:Ekran|Display)\s*:\s*([^\n\r]+)',
            "GPU": r'(?:VGA|GPU|Video|Qrafik)\s*:\s*([^\n\r]+)',
            "OS": r'(?:OS|Əməliyyat sistemi)\s*:\s*([^\n\r]+)'
        }
        for key, pattern in regexes.items():
            m = re.search(pattern, text, flags=re.IGNORECASE)
            if m:
                val = m.group(1).strip()
                clean_val = re.sub(r'[^\w\s\-\.\,\(\)\"\/]', '', val)
                data[key] = clean_val[:65] 
        return data, None
    except Exception as e:
        return data, str(e)

def create_final_design(data, price1, price2, credit):
    try:
        tpl_path = get_absolute_path("template.png")
        base_img = Image.open(tpl_path).convert("RGB")
    except:
        return None

    W, H = base_img.size 
    img = Image.new('RGB', (W, H), color=(54, 54, 64))
    img.paste(base_img, (0, 0))
    draw = ImageDraw.Draw(img)

    # Tənzimləmələr
    VERTICAL_OFFSET = 40 
    ICON_X = 50
    TEXT_X = 110
    START_Y = 340 + VERTICAL_OFFSET 
    ROW_GAP = 100
    ICON_SIZE = 45
    FIXED_PRICE_Y = 1055
    PRICE1_CENTER_X = 275 
    PRICE2_CENTER_X = 515
    WHITE = (255, 255, 255)
    GRAY_TEXT = (200, 200, 200)

    bold_font_name = "TitilliumWeb-Bold.ttf"
    font_brand = load_font(bold_font_name, 60)
    font_model = load_font(bold_font_name, 55)
    font_pn    = load_font(bold_font_name, 34)
    font_spec  = load_font(bold_font_name, 38)
    font_price_main = load_font(bold_font_name, 68)
    
    if data["Brand"]:
        w = draw.textlength(data["Brand"], font=font_brand)
        draw.text(((W - w) / 2, 40 + VERTICAL_OFFSET), data["Brand"], font=font_brand, fill=WHITE)
    if data["Model"]:
        w = draw.textlength(data["Model"], font=font_model)
        draw.text(((W - w) / 2, 110 + VERTICAL_OFFSET), data["Model"], font=font_model, fill=WHITE)
    if data["PartNumber"]:
        w = draw.textlength(data["PartNumber"], font=font_pn)
        draw.text(((W - w) / 2, 220 + VERTICAL_OFFSET), data["PartNumber"], font=font_pn, fill=GRAY_TEXT)

    specs = [
        ("cpu", data["CPU"]), ("ram", data["RAM"]), ("ssd", data["SSD"]),
        ("gpu", data["GPU"]), ("screen", data["Screen"]), ("os", data["OS"]),
    ]
    y = START_Y
    for icon_name, text in specs:
        icon_path = get_absolute_path(os.path.join("icons", f"{icon_name}.png"))
        if os.path.exists(icon_path):
            try:
                icon = Image.open(icon_path).convert("RGBA")
                icon = icon.resize((ICON_SIZE, ICON_SIZE), Image.LANCZOS)
                icon = make_white(icon) 
                img.paste(icon, (ICON_X, y), icon)
            except: pass
        if text:
            max_w = W - TEXT_X - 20
            words = text.split()
            lines = []
            current_line = ""
            for word in words:
                test_line = current_line + " " + word if current_line else word
                if draw.textlength(test_line, font=font_spec) < max_w:
                    current_line = test_line
                else:
                    lines.append(current_line)
                    current_line = word
            lines.append(current_line)
            for i, line in enumerate(lines):
                draw.text((TEXT_X, y + (i * 42) - 5), line, font=font_spec, fill=WHITE)
        y += ROW_GAP + 10

    if price1:
        draw.text((PRICE1_CENTER_X, FIXED_PRICE_Y), str(price1), font=font_price_main, fill=WHITE, anchor="mm")
    if price2:
        draw.text((PRICE2_CENTER_X, FIXED_PRICE_Y), str(price2), font=font_price_main, fill=WHITE, anchor="mm")
    if credit:
        draw.text((W/2, 1130), f"Kredit: {credit}", font=font_spec, fill=WHITE, anchor="mm")
    return img

def create_a4_sheet(images_list):
    A4_W, A4_H = 2480, 3508
    a4_img = Image.new('RGB', (A4_W, A4_H), color=(255, 255, 255))
    gap_x = (A4_W - (800 * 2)) // 3
    gap_y = (A4_H - (1150 * 2)) // 3
    positions = [
        (gap_x, gap_y),                         
        (gap_x + 800 + gap_x, gap_y),           
        (gap_x, gap_y + 1150 + gap_y),          
        (gap_x + 800 + gap_x, gap_y + 1150 + gap_y) 
    ]
    for i, img_bytes in enumerate(images_list):
        if i >= 4: break 
        try:
            card = Image.open(io.BytesIO(img_bytes))
            a4_img.paste(card, positions[i])
        except: pass
    return a4_img


# --- NAVİQASİYA (SOL MENYU) ---
st.sidebar.title("🗂️ Əməliyyat Paneli")
secim = st.sidebar.radio(
    "Zəhmət olmasa seçin:",
    ["🏠 Ana Səhifə", "📄 Müqavilə Hazırla", "📋 Protokol Hazırla", "🏷️ Etiket Hazırla"]
)

# =========================================================
# 1. ANA SƏHİFƏ
# =========================================================
if secim == "🏠 Ana Səhifə":
    col_title, col_logo = st.columns([4, 1])
    with col_title:
        st.title("Xoş Gəlmisiniz! 👋")
    with col_logo:
        logo_goster()

    st.info("Bu sistem vasitəsilə şirkət sənədlərini və etiketləri avtomatik hazırlaya bilərsiniz.")
    st.markdown("""
    ### 🚀 İmkanlar:
    1. **📄 Müqavilə:** Avtomatik alğı-satqı müqavilələri.
    2. **📋 Protokol:** Qiymət razılaşma protokolları.
    3. **🏷️ Etiket:** Saytdan link ilə məhsul etiketlərinin (A4 çap) hazırlanması.
    """)

# =========================================================
# 2. MÜQAVİLƏ SİSTEMİ
# =========================================================
elif secim == "📄 Müqavilə Hazırla":
    col_header, col_logo = st.columns([4, 1])
    with col_header:
        st.header("📄 Müqavilə Generatoru")
    with col_logo:
        logo_goster()
    
    current_dir = os.path.dirname(os.path.abspath(__file__))
    sablon_yolu = os.path.join(current_dir, "sablon.docx")

    if not os.path.exists(sablon_yolu):
        st.error("❌ XƏTA: 'sablon.docx' faylı tapılmadı!")
    else:
        with st.form("muqavile_formu"):
            st.subheader("1. Əsas Məlumatlar")
            c1, c2 = st.columns(2)
            with c1:
                nomre = st.text_input("Müqavilə №", placeholder="055")
                tarix_val = st.date_input("İmzalanma Tarixi", value=date.today())
                tarix_str = tarix_val.strftime("%d.%m.%Y")
                bitme_val = st.date_input("Bitmə Tarixi", value=date.today().replace(year=date.today().year + 1))
                bitme_str = bitme_val.strftime("%d.%m.%Y")
            with c2:
                sirket = st.text_input("Şirkət Adı", placeholder="'ABC' MMC")
                musteri = st.text_input("Direktor", placeholder="Ad Soyad")

            st.subheader("2. Bank Rekvizitləri")
            c3, c4 = st.columns(2)
            with c3:
                voen = st.text_input("Şirkətin VÖEN-i")
                hesab = st.text_input("Hesab (H/h)")
                mh = st.text_input("M/h")
            with c4:
                bank = st.text_input("Bankın Adı")
                bank_voen = st.text_input("Bankın VÖEN-i")
                swift = st.text_input("SWIFT")
                bank_kodu = st.text_input("Bank Kodu")

            submit_btn = st.form_submit_button("Müqaviləni Yarat", type="primary")

        if submit_btn:
            doc = Document(sablon_yolu)
            deyisenler = {
                "{{NOMRE}}": nomre, "{{TARIX}}": tarix_str, "{{BITME_TARIXI}}": bitme_str,
                "{{SIRKET}}": sirket, "{{MUSTERI}}": musteri, "{{VOEN}}": voen,
                "{{HESAB}}": hesab, "{{BANK}}": bank, "{{BANK_VOEN}}": bank_voen,
                "{{MH}}": mh, "{{SWIFT}}": swift, "{{KOD}}": bank_kodu
            }
            
            def replace_bold(p):
                for k, v in deyisenler.items():
                    if k in p.text and v:
                        parts = p.text.split(k)
                        p.clear()
                        for i in range(len(parts)-1):
                            r = p.add_run(parts[i])
                            r = p.add_run(str(v))
                            r.bold = True
                        p.add_run(parts[-1])
                    elif k in p.text:
                        p.text = p.text.replace(k, "")

            for p in doc.paragraphs: replace_bold(p)
            for t in doc.tables:
                for r in t.rows:
                    for c in r.cells:
                        for p in c.paragraphs: replace_bold(p)

            bio = BytesIO()
            doc.save(bio)
            st.success("✅ Müqavilə hazırdır!")
            st.download_button("📥 Yüklə", bio.getvalue(), f"Muqavile_{sirket}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

# =========================================================
# 3. PROTOKOL SİSTEMİ
# =========================================================
elif secim == "📋 Protokol Hazırla":
    col_header, col_logo = st.columns([4, 1])
    with col_header:
        st.header("📋 Protokol Generatoru")
    with col_logo:
        logo_goster()

    current_dir = os.path.dirname(os.path.abspath(__file__))
    sablon_protokol_yolu = os.path.join(current_dir, "sablon_protokol.docx")

    if not os.path.exists(sablon_protokol_yolu):
        st.error("❌ XƏTA: 'sablon_protokol.docx' tapılmadı!")
    else:
        col1, col2 = st.columns([1, 2])
        with col1:
            st.subheader("Məlumatlar")
            nomre = st.text_input("Protokol №", value="1")
            t_val = st.date_input("Tarix", value=date.today())
            t_str = t_val.strftime("%d.%m.%Y")
            sirket = st.text_input("Müştəri", value="", placeholder="Şirkət adı...")
            musteri = st.text_input("Direktor", value="", placeholder="Ad Soyad")
            
            with st.expander("Bank Rekvizitləri"):
                voen = st.text_input("VÖEN")
                hesab = st.text_input("Hesab")
                bank = st.text_input("Bank")
                b_voen = st.text_input("Bank VÖEN")
                mh = st.text_input("M/h")
                swift = st.text_input("SWIFT")
                kod = st.text_input("Kod")

        with col2:
            st.subheader("Məhsullar")
            df = pd.DataFrame(columns=["Malın/Xidmətin Adı", "Ölçü Vahidi", "Miqdar", "Qiymət (AZN)"])
            edited_df = st.data_editor(
                df, num_rows="dynamic", use_container_width=True,
                column_config={
                    "Malın/Xidmətin Adı": st.column_config.TextColumn(width="large", required=True),
                    "Ölçü Vahidi": st.column_config.SelectboxColumn(options=["ədəd", "metr", "kq", "litr", "saat", "gün", "ay", "dəst", "xidmət"], required=True),
                    "Miqdar": st.column_config.NumberColumn(min_value=0, step=1, format="%d"),
                    "Qiymət (AZN)": st.column_config.NumberColumn(min_value=0, format="%.2f ₼")
                }
            )
            
            cem = 0.0
            if not edited_df.empty:
                edited_df["Miqdar"] = pd.to_numeric(edited_df["Miqdar"], errors='coerce').fillna(0)
                edited_df["Qiymət (AZN)"] = pd.to_numeric(edited_df["Qiymət (AZN)"], errors='coerce').fillna(0)
                edited_df["Məbləğ"] = edited_df["Miqdar"] * edited_df["Qiymət (AZN)"]
                cem = edited_df["Məbləğ"].sum()
            
            edv = cem * 0.18
            yekun = cem + edv
            
            c_a, c_b, c_c = st.columns(3)
            c_a.metric("Cəmi", f"{cem:.2f} ₼")
            c_b.metric("ƏDV (18%)", f"{edv:.2f} ₼")
            c_c.metric("YEKUN", f"{yekun:.2f} ₼")

        if st.button("Protokolu Hazırla", type="primary"):
            doc = Document(sablon_protokol_yolu)
            style = doc.styles['Normal']
            style.font.name = 'Calibri'
            style.font.size = Pt(11)

            deyisenler = {
                "{{NOMRE}}": nomre, "{{TARIX}}": t_str, "{{SIRKET}}": sirket,
                "{{MUSTERI}}": musteri, "{{VOEN}}": voen, "{{HESAB}}": hesab,
                "{{BANK}}": bank, "{{BANK_VOEN}}": b_voen, "{{MH}}": mh,
                "{{SWIFT}}": swift, "{{KOD}}": kod
            }

            def replace_bold_calibri(paragraph):
                for k, v in deyisenler.items():
                    if k in paragraph.text:
                        if v:
                            parts = paragraph.text.split(k)
                            paragraph.clear()
                            for i in range(len(parts)-1):
                                r = paragraph.add_run(parts[i])
                                r.font.name = 'Calibri'
                                r = paragraph.add_run(str(v))
                                r.bold = True
                                r.font.name = 'Calibri'
                            r = paragraph.add_run(parts[-1])
                            r.font.name = 'Calibri'
                        else:
                            paragraph.text = paragraph.text.replace(k, "")

            for p in doc.paragraphs: replace_bold_calibri(p)
            for t in doc.tables:
                for r in t.rows:
                    for c in r.cells:
                        for p in c.paragraphs: replace_bold_calibri(p)

            try:
                table = doc.tables[0]
                for i in range(len(table.rows)-1, 0, -1):
                    table.rows[i]._element.getparent().remove(table.rows[i]._element)
                
                for idx, row in edited_df.iterrows():
                    nr = table.add_row()
                    vals = [str(idx+1), str(row["Malın/Xidmətin Adı"] or ""), str(row["Ölçü Vahidi"] or ""),
                            str(int(row["Miqdar"])), f"{row['Qiymət (AZN)']:.2f}", f"{row['Məbləğ']:.2f}"]
                    for i, val in enumerate(vals):
                        nr.cells[i].text = val
                        for p in nr.cells[i].paragraphs:
                            for r in p.runs: r.font.name = 'Calibri'
                
                def add_sum(l, v):
                    r = table.add_row()
                    r.cells[4].text = l
                    r.cells[5].text = v
                    for c in [r.cells[4], r.cells[5]]:
                        for p in c.paragraphs:
                            for run in p.runs:
                                run.font.name = 'Calibri'
                                run.bold = True
                
                add_sum("Cəmi", f"{cem:.2f}")
                add_sum("ƏDV 18%", f"{edv:.2f}")
                add_sum("Yekun", f"{yekun:.2f}")

            except Exception as e:
                st.error(f"Cədvəl xətası: {e}")

            bio = BytesIO()
            doc.save(bio)
            st.success("✅ Protokol hazırdır!")
            st.download_button("📥 Yüklə", bio.getvalue(), f"Protokol_{sirket}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

# =========================================================
# 4. ETİKET SİSTEMİ (YENİ ƏLAVƏ)
# =========================================================
elif secim == "🏷️ Etiket Hazırla":
    col_header, col_logo = st.columns([4, 1])
    with col_header:
        st.header("🏷️ LaptopMarket Etiket Generatoru")
    with col_logo:
        logo_goster()

    # Sidebar-da Siyahı Görünüşü
    st.sidebar.markdown("---")
    st.sidebar.header(f"🖨️ Çap Siyahısı: {len(st.session_state.print_queue)} / 4")
    
    if st.sidebar.button("🗑️ Siyahını Təmizlə"):
        st.session_state.print_queue = []
        st.rerun()

    if len(st.session_state.print_queue) > 0:
        st.sidebar.success("Hazırdır!")
        if st.sidebar.button("📄 A4 Yarat və Yüklə"):
            final_a4 = create_a4_sheet(st.session_state.print_queue)
            st.write("## 🎉 Çap üçün A4 Vərəqi")
            st.image(final_a4, caption="Önizləmə", use_container_width=True)
            buf = io.BytesIO()
            final_a4.save(buf, format="PDF") 
            st.download_button("📥 A4 PDF Yüklə", buf.getvalue(), "etiketler_A4.pdf", "application/pdf")

    # Əsas Etiket Formu
    col1, col2 = st.columns([3, 1])
    url = col1.text_input("Məhsul Linki:", placeholder="Sayt linkini yapışdırın...")
    if col2.button("Məlumatı Gətir 🔍"):
        with st.spinner("Gətirilir..."):
            data, err = scrape_smart(url)
            if not err:
                st.session_state.specs = data
                st.success("Məlumatlar tapıldı!")
            else:
                st.error("Məlumat tapılmadı, zəhmət olmasa əllə yazın.")

    st.write("### Etiket Dizaynı")
    with st.form("design_form"):
        s = st.session_state.specs
        
        c1, c2, c3 = st.columns([2, 2, 2])
        brand = c1.text_input("Brend", s["Brand"])
        model = c2.text_input("Model", s["Model"])
        pn    = c3.text_input("Part No", s["PartNumber"])
        
        st.markdown("---")
        c4, c5 = st.columns(2)
        cpu = c4.text_input("CPU", s["CPU"])
        ram = c5.text_input("RAM", s["RAM"])
        ssd = c4.text_input("SSD", s["SSD"])
        gpu = c5.text_input("GPU", s["GPU"])
        scr = c4.text_input("Ekran", s["Screen"])
        os_ = c5.text_input("OS", s["OS"])
        
        st.markdown("---")
        st.write("💰 **Qiymət Xanaları**")
        cp1, cp2, cc = st.columns(3)
        
        p1_val = cp1.text_input("Sol Xana (Köhnə/Qırmızı)", value=s.get("PriceOld", ""))
        p2_val = cp2.text_input("Orta Xana (Yeni/Yaşıl)", value=s.get("PriceNew", ""))
        cred_val = cc.text_input("Kredit", placeholder="65 AZN/ay")
        
        submitted = st.form_submit_button("🎨 Etiketi Hazırla")
        
        if submitted:
            final_data = {
                "Brand": brand, "Model": model, "PartNumber": pn,
                "CPU": cpu, "RAM": ram, "SSD": ssd, 
                "GPU": gpu, "Screen": scr, "OS": os_
            }
            img = create_final_design(final_data, p1_val, p2_val, cred_val)
            st.session_state.generated_image = img

    # Nəticə Göstərilməsi
    if st.session_state.generated_image:
        st.write("---")
        c_img, c_btn = st.columns([2, 1])
        c_img.image(st.session_state.generated_image, caption="Hazır Etiket", width=300)
        c_btn.write("### Bəyəndiniz?")
        if c_btn.button("➕ Siyahıya Əlavə Et (Çap Üçün)"):
            buf = io.BytesIO()
            st.session_state.generated_image.save(buf, format="PNG")
            st.session_state.print_queue.append(buf.getvalue())
            st.success(f"Əlavə olundu! Siyahıda {len(st.session_state.print_queue)} ədəd var.")
            st.rerun()